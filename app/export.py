"""统计结果导出为 Excel。

工作簿的组装写成纯函数（不读数据库、不依赖 Flask），路由只负责取数据与下发，
这样导出格式可以单独测试，统计口径也继续复用 stats.question_stats。
"""

from __future__ import annotations

import re
from datetime import datetime, timezone
from io import BytesIO
from typing import Any

from flask import Blueprint, abort, send_file
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.worksheet.worksheet import Worksheet

from .stats import question_stats
from .utils import get_db, load_questions, login_required, owned_survey_or_404, to_object_id

bp = Blueprint("export", __name__)

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

STATUS_TEXT = {"draft": "草稿", "published": "已发布", "closed": "已关闭"}
TYPE_TEXT = {"single": "单选题", "multiple": "多选题", "text": "文本填空题", "number": "数字填空题"}

HEADERS = ["题号", "题型", "题干", "选项 / 取值", "人数", "占比", "统计值"]
COLUMN_WIDTHS = [8, 11, 34, 46, 10, 10, 18]

HEADER_FILL = PatternFill("solid", fgColor="E8EFF9")
TITLE_FONT = Font(name="微软雅黑", size=11, bold=True)
BODY_FONT = Font(name="微软雅黑", size=10)
THIN = Side(style="thin", color="D5DBE5")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _local(moment: datetime | None) -> datetime | None:
    """MongoDB 取出来的是带时区的 UTC，Excel 需要不带时区的本地时间。"""
    if moment is None:
        return None
    if moment.tzinfo is not None:
        moment = moment.astimezone()
    return moment.replace(tzinfo=None)


def _style_header(ws: Worksheet, row: int, columns: int) -> None:
    for col in range(1, columns + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = TITLE_FONT
        cell.fill = HEADER_FILL
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center")


def _style_body(ws: Worksheet, first_row: int) -> None:
    for row in ws.iter_rows(min_row=first_row, max_row=ws.max_row, max_col=len(HEADERS)):
        for cell in row:
            cell.font = BODY_FONT
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=cell.column in (3, 4))


def _build_info_sheet(wb: Workbook, survey: dict, total: int, scope: str) -> None:
    ws = wb.active
    ws.title = "问卷信息"
    ws.append(["项目", "值"])
    _style_header(ws, 1, 2)
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 60

    rows: list[tuple[str, Any]] = [
        ("问卷标题", survey.get("title", "")),
        ("问卷说明", survey.get("description") or ""),
        ("状态", STATUS_TEXT.get(survey.get("status"), survey.get("status", ""))),
        ("创建时间", _local(survey.get("created_at"))),
        ("发布时间", _local(survey.get("published_at")) or "—"),
        ("截止时间", _local(survey.get("deadline")) or "不限"),
        ("题目数量", survey.get("question_count", 0)),
        ("回答数量", total),
        ("匿名填写", "是" if survey.get("allow_anonymous") else "否"),
        ("允许重复填写", "是" if survey.get("allow_multiple") else "否"),
        ("导出范围", scope),
        ("导出时间", datetime.now().replace(microsecond=0)),
    ]
    for label, value in rows:
        ws.append([label, value])

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=2):
        for cell in row:
            cell.font = BODY_FONT
            cell.border = BORDER
            cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=cell.column == 2)
        row[0].font = TITLE_FONT

    for row in range(2, ws.max_row + 1):
        cell = ws.cell(row=row, column=2)
        if isinstance(cell.value, datetime):
            cell.number_format = "yyyy-mm-dd hh:mm"


def _write_question(ws: Worksheet, number: int, question: dict, data: dict) -> None:
    qtype = question.get("type")
    title = question.get("title", "")
    type_text = TYPE_TEXT.get(qtype, qtype)

    if qtype in ("single", "multiple"):
        for option in data.get("options", []):
            ws.append(
                [
                    number,
                    type_text,
                    title,
                    f"{option['key']}. {option['label']}",
                    option["count"],
                    round(option["percent"] / 100, 4),
                    "",
                ]
            )
        return

    if qtype == "number":
        ws.append([number, type_text, title, "平均值", "", "", data.get("average")])
        ws.append([number, type_text, title, "最小值", "", "", data.get("minimum")])
        ws.append([number, type_text, title, "最大值", "", "", data.get("maximum")])
        ws.append([number, type_text, title, "样本数", "", "", len(data.get("values") or [])])
        return

    # 文本题：先给人数与平均长度，再逐条列出内容
    average = data.get("average_length")
    ws.append(
        [
            number,
            type_text,
            title,
            "填写人数 / 平均长度",
            data.get("answered", 0),
            "",
            f"{average} 字" if average is not None else "—",
        ]
    )
    entries = data.get("entries") or []
    if entries:
        ws.append([number, type_text, title, "全部填写内容：", "", "", ""])
        for entry in entries:
            ws.append([number, type_text, title, entry.get("value", ""), "", "", ""])


def build_workbook(
    survey: dict,
    rows: list[tuple[dict, dict]],
    total_responses: int,
    scope: str = "整卷",
) -> Workbook:
    """把问卷、题目与统计结果组装成工作簿。rows 为 [(question, stats), ...]。"""
    wb = Workbook()
    _build_info_sheet(wb, survey, total_responses, scope)

    ws = wb.create_sheet("统计汇总")
    ws.append(HEADERS)
    _style_header(ws, 1, len(HEADERS))
    for index, width in enumerate(COLUMN_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width

    if total_responses == 0:
        # 合并单元格只保留左上角的值，所以提示文字必须写在 A 列
        ws.append(["本问卷暂无回答，以下题目统计均为 0。", "", "", "", "", "", ""])
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(HEADERS))
        ws.cell(row=2, column=1).alignment = Alignment(horizontal="center", vertical="center")

    for number, (question, data) in enumerate(rows, start=1):
        _write_question(ws, number, question, data)

    _style_body(ws, 2)
    for row in range(2, ws.max_row + 1):
        percent = ws.cell(row=row, column=6)
        if isinstance(percent.value, (int, float)):
            percent.number_format = "0.0%"
            percent.alignment = Alignment(horizontal="right", vertical="top")
        for col in (1, 5, 7):
            cell = ws.cell(row=row, column=col)
            if isinstance(cell.value, (int, float)):
                cell.alignment = Alignment(horizontal="right", vertical="top")

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{ws.max_row}"

    # 打印时横向铺满一页，并在每页重复表头
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.print_title_rows = "1:1"

    wb.properties.title = f"{survey.get('title', '问卷')} · 统计导出"
    wb.properties.creator = "问卷系统"
    return wb


def safe_filename(name: str, limit: int = 80) -> str:
    cleaned = re.sub(r'[\\/:*?"<>|\r\n\t]', "_", name).strip(" .")
    return (cleaned or "统计")[:limit]


def _send(wb: Workbook, survey: dict):
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    stamp = datetime.now().strftime("%Y%m%d")
    filename = safe_filename(f"{survey.get('title', '问卷')}-统计-{stamp}.xlsx")
    return send_file(
        buffer,
        as_attachment=True,
        download_name=filename,
        mimetype=XLSX_MIME,
    )


@bp.get("/surveys/<sid>/stats/export.xlsx")
@login_required
def export_survey(sid: str):
    survey = owned_survey_or_404(sid)
    questions = load_questions(survey["_id"])
    rows = [(question, question_stats(survey, question)) for question in questions]
    total = get_db().responses.count_documents({"survey_id": survey["_id"]})
    wb = build_workbook(survey, rows, total, scope="整卷")
    return _send(wb, survey)


@bp.get("/surveys/<sid>/stats/<qid>/export.xlsx")
@login_required
def export_question(sid: str, qid: str):
    survey = owned_survey_or_404(sid)
    oid = to_object_id(qid)
    question = get_db().questions.find_one({"_id": oid, "survey_id": survey["_id"]}) if oid else None
    if question is None:
        abort(404)
    total = get_db().responses.count_documents({"survey_id": survey["_id"]})
    scope = f"第 {question.get('order', 0) + 1} 题 · {question.get('title', '')}"
    wb = build_workbook(survey, [(question, question_stats(survey, question))], total, scope=scope)
    return _send(wb, survey)
