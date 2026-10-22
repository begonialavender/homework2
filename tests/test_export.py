"""Excel 导出测试：文件能否下载、内容与统计口径是否一致、权限是否守住。"""

from __future__ import annotations

from io import BytesIO

from openpyxl import load_workbook

from app.export import safe_filename
from tests.helpers import (
    add_multiple, add_number, add_single, add_text,
    create_survey, login, publish, register, submit,
)


def build(client, db, title="课程反馈"):
    register(client, "alice")
    login(client, "alice")
    survey = create_survey(client, db, title=title)
    q1 = add_single(client, db, survey, title="性别", options="男\n女\n其他")
    q2 = add_multiple(client, db, survey, title="水果", options="苹果\n香蕉\n西瓜\n葡萄")
    q3 = add_number(client, db, survey, title="年龄", min_value=0, max_value=120, integer=True)
    q4 = add_text(client, db, survey, title="建议")
    publish(client, survey)

    rows = [
        {str(q1["_id"]): "A", str(q2["_id"]): ["A", "B"], str(q3["_id"]): "10", str(q4["_id"]): "hi"},
        {str(q1["_id"]): "A", str(q2["_id"]): ["A"], str(q3["_id"]): "20", str(q4["_id"]): "hello"},
        {str(q1["_id"]): "B", str(q2["_id"]): ["B", "C"], str(q3["_id"]): "30", str(q4["_id"]): "abc"},
    ]
    for row in rows:
        assert submit(client, survey["slug"], row).status_code == 200
    return survey, q1, q2, q3, q4


def load(resp):
    assert resp.status_code == 200
    return load_workbook(BytesIO(resp.data))


def sheet_rows(ws):
    return list(ws.iter_rows(values_only=True))


def find_row(rows, column, text):
    for row in rows:
        if row[column] == text:
            return row
    return None


class TestDownload:
    def test_full_export_is_xlsx(self, client, db):
        survey, *_ = build(client, db)
        resp = client.get(f"/surveys/{survey['_id']}/stats/export.xlsx")
        assert resp.status_code == 200
        assert resp.mimetype.endswith("spreadsheetml.sheet")
        assert ".xlsx" in resp.headers["Content-Disposition"]
        assert len(resp.data) > 4000

    def test_single_question_export(self, client, db):
        survey, q1, *_ = build(client, db)
        resp = client.get(f"/surveys/{survey['_id']}/stats/{q1['_id']}/export.xlsx")
        assert resp.status_code == 200
        assert ".xlsx" in resp.headers["Content-Disposition"]

    def test_no_responses_still_exports(self, client, db):
        register(client, "bob")
        login(client, "bob")
        survey = create_survey(client, db)
        add_single(client, db, survey)
        publish(client, survey)

        wb = load(client.get(f"/surveys/{survey['_id']}/stats/export.xlsx"))
        rows = sheet_rows(wb["统计汇总"])
        assert any("暂无回答" in str(cell) for row in rows for cell in row)

    def test_export_requires_login(self, client, db):
        survey, *_ = build(client, db)
        client.post("/logout")
        resp = client.get(f"/surveys/{survey['_id']}/stats/export.xlsx")
        assert resp.status_code == 302
        assert "/login" in resp.headers["Location"]

    def test_other_user_cannot_export(self, client, db):
        survey, *_ = build(client, db)
        client.post("/logout")
        register(client, "bob")
        login(client, "bob")
        assert client.get(f"/surveys/{survey['_id']}/stats/export.xlsx").status_code == 404

    def test_unknown_question_404(self, client, db):
        survey, *_ = build(client, db)
        resp = client.get(f"/surveys/{survey['_id']}/stats/5f0000000000000000000000/export.xlsx")
        assert resp.status_code == 404


class TestWorkbook:
    def test_sheets_and_info(self, client, db):
        survey, *_ = build(client, db)
        wb = load(client.get(f"/surveys/{survey['_id']}/stats/export.xlsx"))

        assert wb.sheetnames == ["问卷信息", "统计汇总"]
        info = dict(sheet_rows(wb["问卷信息"]))
        assert info["问卷标题"] == "课程反馈"
        assert info["状态"] == "已发布"
        assert info["回答数量"] == 3
        assert info["题目数量"] == 4
        assert info["导出范围"] == "整卷"

    def test_single_choice_counts(self, client, db):
        survey, *_ = build(client, db)
        rows = sheet_rows(load(client.get(f"/surveys/{survey['_id']}/stats/export.xlsx"))["统计汇总"])

        assert rows[0][:7] == ("题号", "题型", "题干", "选项 / 取值", "人数", "占比", "统计值")
        assert find_row(rows, 3, "A. 男")[4] == 2
        assert find_row(rows, 3, "B. 女")[4] == 1
        assert find_row(rows, 3, "C. 其他")[4] == 0

    def test_percent_is_numeric(self, client, db):
        survey, *_ = build(client, db)
        wb = load(client.get(f"/surveys/{survey['_id']}/stats/export.xlsx"))
        ws = wb["统计汇总"]

        for row in ws.iter_rows(min_row=2):
            if row[3].value == "A. 男":
                assert abs(row[5].value - 2 / 3) < 0.01
                assert row[5].number_format == "0.0%"
                break
        else:  # pragma: no cover
            raise AssertionError("没有找到 A. 男 这一行")

    def test_multiple_choice_counts(self, client, db):
        survey, *_ = build(client, db)
        rows = sheet_rows(load(client.get(f"/surveys/{survey['_id']}/stats/export.xlsx"))["统计汇总"])
        assert find_row(rows, 3, "A. 苹果")[4] == 2
        assert find_row(rows, 3, "D. 葡萄")[4] == 0

    def test_number_summary(self, client, db):
        survey, *_ = build(client, db)
        rows = sheet_rows(load(client.get(f"/surveys/{survey['_id']}/stats/export.xlsx"))["统计汇总"])
        assert find_row(rows, 3, "平均值")[6] == 20
        assert find_row(rows, 3, "最小值")[6] == 10
        assert find_row(rows, 3, "最大值")[6] == 30
        assert find_row(rows, 3, "样本数")[6] == 3

    def test_text_entries_listed(self, client, db):
        survey, *_ = build(client, db)
        rows = sheet_rows(load(client.get(f"/surveys/{survey['_id']}/stats/export.xlsx"))["统计汇总"])
        values = {row[3] for row in rows}
        assert {"hi", "hello", "abc"} <= values

    def test_single_question_scope(self, client, db):
        survey, q1, q2, *_ = build(client, db)
        wb = load(client.get(f"/surveys/{survey['_id']}/stats/{q1['_id']}/export.xlsx"))
        rows = sheet_rows(wb["统计汇总"])
        titles = {row[2] for row in rows[1:]}
        assert titles == {"性别"}
        assert q2["title"] not in titles

        info = dict(sheet_rows(wb["问卷信息"]))
        assert info["导出范围"].startswith("第 1 题")

    def test_freeze_and_filter(self, client, db):
        survey, *_ = build(client, db)
        ws = load(client.get(f"/surveys/{survey['_id']}/stats/export.xlsx"))["统计汇总"]
        assert ws.freeze_panes == "A2"
        assert ws.auto_filter.ref.startswith("A1:G")
        assert ws.page_setup.orientation == "landscape"
        assert ws.sheet_properties.pageSetUpPr.fitToPage is True


class TestFilename:
    def test_illegal_characters_removed(self):
        assert safe_filename('a/b:c*?"<>|d.xlsx') == "a_b_c______d.xlsx"

    def test_empty_falls_back(self):
        assert safe_filename("   ") == "统计"

    def test_length_limited(self):
        assert len(safe_filename("长" * 200)) == 80
